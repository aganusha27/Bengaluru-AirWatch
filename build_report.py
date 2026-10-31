"""Create an evidence-driven DOCX report and supporting handoff files after main.py."""
import json
import re
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parent
PAPER='Temporal dynamics of urban air pollutants and their correlation with associated meteorological parameters: an investigation in northern Indian cities'
REFERENCES=[
 ('Soni et al. 2024',PAPER+'. Environmental Monitoring and Assessment 196, article 505.','https://doi.org/10.1007/s10661-024-12678-z'),
 ('Publisher record','Verified title, publication date and abstract.','https://link.springer.com/article/10.1007/s10661-024-12678-z'),
 ('PubMed','Bibliographic and abstract fallback, PMID 38700603.','https://pubmed.ncbi.nlm.nih.gov/38700603/'),
 ('Author supplied full text','Methodology, R version, study period inconsistencies and limitations. The author-hosted page restricts some automated access; publisher and PubMed links provide bibliographic fallbacks.','https://www.researchgate.net/publication/380321341_Temporal_dynamics_of_urban_air_pollutants_and_their_correlation_with_associated_meteorological_parameters_an_investigation_in_northern_Indian_cities'),
 ('OpenCity and KSPCB','Bengaluru Hourly Air Quality Reports. Selected 2024-2025 station resources; see data/manifest.json for exact downloads.','https://data.opencity.in/dataset/bengaluru-hourly-air-quality-reports'),
 ('Apache Spark 4.0.1','Installation, Java and Python requirements.','https://spark.apache.org/docs/4.0.1/api/python/getting_started/install.html'),
 ('Apache Spark 4.0.1','Exact percentile aggregation used for the spike threshold.','https://spark.apache.org/docs/4.0.1/api/python/reference/pyspark.sql/api/pyspark.sql.functions.percentile.html')]

def build():
    def load(name):return pd.read_csv(ROOT/f'outputs/tables/{name}.csv')
    m=json.loads((ROOT/'outputs/metrics/performance.json').read_text())
    check=json.loads((ROOT/'outputs/metrics/verification.json').read_text());assert check['status']=='PASS'
    manifest=json.loads((ROOT/'data/manifest.json').read_text())
    stats=load('station_statistics').sort_values('mean',ascending=False)
    common=load('common_case_station_statistics').sort_values('mean',ascending=False)
    spikes=load('spike_summary');pairs=load('synchronization_pairs').sort_values('lift',ascending=False)
    inventory=load('dataset_inventory');audit=load('cleaning_audit');coverage=load('hourly_coverage')
    complete=load('complete_case_spike_distribution');ncomplete=int(complete['count'].sum())
    multiple=int(complete.loc[complete.spiking_stations>=2,'count'].sum())
    one=int(complete.loc[complete.spiking_stations==1,'count'].sum())
    nspiking=int(complete.loc[complete.spiking_stations>0,'count'].sum())
    high,low=stats.iloc[0],stats.iloc[-1];top_pair=pairs.iloc[0]
    def short(x):return str(x).split(',')[0]
    def f(x,n=2):return f'{x:,.{n}f}'
    doc=Document();sec=doc.sections[0]
    # Clear Word template paragraph borders, including the default Title rule.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.78)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.12
    for name,size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(15)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Bengaluru AirWatch  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    md=[];evidence=[]
    def h(t,level=1):doc.add_heading(t,level);md.append('#'*(level+1)+' '+t+'\n')
    def p(t):doc.add_paragraph(t);md.append(t+'\n')
    def bullet(t):doc.add_paragraph(t,style='List Bullet');md.append('- '+t)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,value in zip(t.rows[0].cells,headers):c.text=str(value)
        trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for c,v in zip(cells,row):c.text=str(v)
        for row in t.rows:
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.space_after=Pt(3)
                    for r in pp.runs:r.font.size=Pt(9)
        md.extend([' | '.join(headers),' | '.join(['---']*len(headers))]+[' | '.join(map(str,row)) for row in rows]+[''])
    def hyperlink(label,url):
        pp=doc.add_paragraph();rid=pp.part.relate_to(url,RT.HYPERLINK,is_external=True)
        link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
        run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'155A92');pr.append(color);run.append(pr)
        text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);pp._p.append(link)
        md.append(f'[{label}]({url})\n')
    def claim(name,value,source):evidence.append(dict(claim=name,value=value,source=source))

    doc.add_paragraph('Bengaluru AirWatch Air Pollution Analysis','Title')
    p('Station patterns and synchronized PM2.5 spikes using Apache Spark')
    p('Ramaiah Institute of Technology\nDepartment of Computer Science and Engineering\nBig Data Analytics\nSubject code CSE557\nAcademic term 2026')
    p('Historical observations from six Bengaluru monitoring stations\nStudy window 1 January 2024 to 31 December 2025 in IST')
    team_path=ROOT/'report/team_details.json'
    team=json.loads(team_path.read_text()) if team_path.exists() else {}
    if team.get('members'):
        if any(v.get('responsibility') for v in team['members']):
            table(['Student name','USN','Responsibility'],[[v.get('name',''),v.get('usn',''),v.get('responsibility','')] for v in team['members']])
        else:
            table(['Student name','USN'],[[v.get('name',''),v.get('usn','')] for v in team['members']])
    if team.get('faculty_guide'):p('Faculty guide: '+team['faculty_guide'])
    p('Analysis executed on '+m['run_time_ist'][:10]+'. All reported project measurements come from the saved execution outputs.')
    doc.add_page_break()
    h('Report contents')
    sections=[
      'Real world scenario and problem statement','Recent research paper and existing approach',
      'Objectives and proposed analysis','Dataset identification and preparation',
      'Big Data implementation and workflow','Results visualization and paper comparison',
      'Performance evaluation and observations','Tools responsible use and references']
    for i,name in enumerate(sections,1):p(f'{i}  {name}')
    p('The eight main sections follow the faculty assignment. Subsections include the abstract, research review, implementation, empirical results, performance, limitations and execution appendix.')
    h('Abstract',2)
    p(f'This project analyzes {m["raw_records"]:,} source rows from six Bengaluru air-quality CSV files. After cleaning and hourly coverage checks, {m["output_hourly_records"]:,} station-hour PM2.5 observations remain. {short(high.station)} has the highest observed mean at {f(high["mean"])} µg/m³, compared with {f(low["mean"])} µg/m³ at {short(low.station)}. Station-specific 95th-percentile thresholds identify unusually high hours. Of {ncomplete:,} hours observed at all six stations, {multiple:,} contain spikes at two or more stations. Pairwise coincidence is compared with a marginal-probability independence baseline. Eight charts and the analytical tables are generated using real data. Independent pandas recomputation verifies the Spark outputs. These are descriptive historical findings; they establish neither pollutant transport nor prediction performance.')
    claim('Raw source rows',m['raw_records'],'data/manifest.json records sum')
    claim('Valid station hours',m['output_hourly_records'],'outputs/metrics/performance.json output_hourly_records')
    claim('Complete six station hours',ncomplete,'outputs/tables/complete_case_spike_distribution.csv count sum')
    claim('Complete hours with multiple spikes',multiple,'outputs/tables/complete_case_spike_distribution.csv spiking_stations >= 2')
    doc.add_page_break()
    h('1 Real world scenario and problem statement')
    h('Introduction',2)
    p('A citywide average can hide differences between monitoring locations. Bengaluru AirWatch examines how measured PM2.5 varies between stations, across hours and across months, and whether unusually high readings coincide. It uses historical monitoring data rather than live streaming. The analysis supports comparison and investigation, rather than issuing health advisories or identifying emission sources.')
    h('Real world scenario and problem statement',2)
    p('The central question is whether station-relative PM2.5 spikes occur together across Bengaluru or remain specific to individual locations. The practical difficulty is aligning heterogeneous coverage without turning missing values into clean-air observations. A reproducible DataFrame workflow is useful for ingestion, quality auditing, aggregation and timestamp joins. This local dataset demonstrates these operations but does not require a distributed cluster.')
    h('Stakeholders and analytical questions',2)
    table(['Stakeholder','Question supported by this analysis'],[
      ['Pollution monitoring staff','Which stations have higher observed levels and fluctuations?'],
      ['City researchers','Which hours and months show higher concentrations?'],
      ['Students and project reviewers','Can counts, thresholds and comparisons be reproduced from source files?'],
      ['Urban planning analysts','Which station pairs show coincident high readings that merit further investigation?']])
    h('2 Recent research paper and existing approach')
    h('Research paper review',2)
    p('Soni and colleagues published '+PAPER+' in Environmental Monitoring and Assessment, volume 196, article 505, on 3 May 2024 [1]. It investigates six pollutants at ten northern Indian monitoring locations and their meteorological associations. The publisher abstract reports spatial differences and positive PM2.5–NO2 associations [2].')
    p('The author-supplied text describes CPCB observations, temporal summaries, bivariate polar plots, HYSPLIT trajectories and analysis in R 4.2.0 [4]. Its period descriptions differ: the overview and figures extend from January 2017 to March 2021, while parts of the methodology refer to March or December 2020. This discrepancy is retained rather than resolved by assumption.')
    h('Existing approach and limitations',2)
    p('The reviewed work connects concentrations with meteorology and transport indicators. It does not supply the present project’s Spark pipeline or station-pair spike baseline. Its northern locations and earlier period limit direct transfer to Bengaluru. Reported data gaps and inconsistent period descriptions also warrant caution. Our comparison therefore concerns analytical scope and workflow; it is not a numerical replication or a causal source-attribution study.')
    h('3 Objectives and proposed analysis')
    h('Project objectives',2)
    for text in [
      'Prepare six source files and account for every input row, missing PM2.5 value, duplicate and rejected hourly bucket; preserve immutable raw data and publish file checksums.',
      'Implement and execute Spark DataFrame and SQL analyses producing station summaries, five temporal tables, 95th-percentile spike detection and all 15 station-pair synchronization comparisons.',
      'Evaluate the analysis using independent pandas recomputation, eight meaningful visualizations, measured execution times and a three-repeat warm aggregation comparison on the same hourly data.']:
        bullet(text)
    h('Proposed methodology',2)
    p('Parse source timestamps with their encoded +0000 offset and express calendar calculations in Asia/Kolkata. Restrict observations to the local 2024–2025 window. Reject absent station identifiers, invalid timestamps, missing or nonnumeric PM2.5, negative or nonfinite concentrations and timestamps outside the quarter-hour grid. Keep zero if measured. Collapse identical selected observations and average conflicting values at the same station and instant. Compute hourly means only when at least three distinct quarter-hour observations remain.')
    p('For each station, compute the exact interpolated 95th percentile of its valid hourly means. An hour is a spike only when its mean is strictly greater than that threshold. Ties can make the spike rate differ from 5%. Thresholds describe relative extremes, not regulatory exceedances. They are estimated on the analyzed period and are not evaluated as an out-of-sample detector.')
    p('For a pair with N overlapping valid hours, let A and B be its individual spike counts on those same hours and J the simultaneous spike count. The independence expectation is A × B / N; lift is J divided by that expectation. Jaccard overlap is J / (A + B − J). These descriptive baselines do not remove seasonality or serial dependence and are not significance tests. A second table uses only hours observed at all six stations to summarize citywide coincidence.')
    h('4 Dataset identification and preparation')
    h('Dataset identification and description',2)
    p(f'The source is the OpenCity Bengaluru Hourly Air Quality Reports catalog [5], under its KSPCB organization. Six selected resources are labelled 15-minute AQI Data for 2024–25, including two CPCB-labelled stations. Their actual schema contains concentrations and weather variables, with no AQI column. PM2.5 is measured in µg/m³ and is the sole analysis variable. CO uses mg/m³ and NOx uses ppb; they are not mixed with PM2.5. Total source size is {m["raw_bytes"]:,} bytes ({m["raw_bytes"]/1048576:.2f} MiB). All six schemas contain {int(inventory['columns'].iloc[0])} matching columns.')
    rows=[]
    for r in inventory.itertuples():rows.append([short(r.station),f'{r.records:,}',f'{r.bytes/1048576:.2f}',r.start_utc[:10],r.end_utc[:10]])
    table(['Station','Source rows','MiB','Start UTC','End UTC'],rows)
    p('Timestamps explicitly encode UTC. The project honors that encoding and converts to IST; the original instrument timezone convention has not been independently confirmed with the publisher. The 2026 local-time spillover caused by conversion is excluded. Resource filenames and exact URLs, byte sizes, SHA-256 checksums and first and last raw records are preserved in data/manifest.json. Catalog metadata is preserved in data/catalog_snapshot.json.')
    h('Dataset preprocessing',2)
    c=m['cleaning_counts'];rejected=c['raw']-c['valid']
    table(['Processing stage','Rows or buckets'],[
      ['Source CSV rows',f'{c["raw"]:,}'],['Rejected source rows',f'{rejected:,}'],['Valid source observations',f'{c["valid"]:,}'],
      ['After identical selected observations collapse',f'{c["distinct"]:,}'],['Unique station and timestamp observations',f'{c["deduplicated"]:,}'],
      ['Candidate station-hour buckets',f'{c["candidate_hours"]:,}'],['Hours rejected for fewer than three quarters',f'{c["candidate_hours"]-c["hourly"]:,}'],
      ['Retained station-hour means',f'{c["hourly"]:,}']])
    reasons=audit.groupby('reason')['count'].sum()
    for reason,count in reasons.items():bullet(f'{reason}: {int(count):,} source rows.')
    p(f'The files contain six single-field year separator rows. They have missing station and timestamp fields and are included in the rejection audit. Identical selected observations removed: {c["valid"]-c["distinct"]:,}; additional conflicting station-timestamp values combined: {c["distinct"]-c["deduplicated"]:,}. No missing measurement is imputed. Valid high concentrations are retained because no verified sensor-specific upper bound is available. This can leave instrument outliers in the maxima and percentile estimates. Quality counts refer to PM2.5 processing; other variables remain unchanged in raw files.')
    claim('Rejected source rows',rejected,'performance.json cleaning_counts raw minus valid')
    h('Big Data characteristics',2)
    table(['Characteristic','Evidence and interpretation'],[
      ['Volume',f'{m["raw_records"]:,} rows and {m["raw_bytes"]/1048576:.2f} MiB across six files; manageable locally, not evidence of very large scale.'],
      ['Velocity','Source resources are labelled 15-minute observations; the submitted workflow is historical batch processing.'],
      ['Variety','Station metadata, offset-aware timestamps, pollutant concentrations and meteorological fields; matching CSV schemas in the selected subset.'],
      ['Veracity','Missing measurements, separator rows, incomplete station coverage, repeated readings and uncertain sensor validation require explicit audits.']])
    h('5 Big Data implementation and workflow')
    h('System architecture and workflow',2)
    for text in ['OpenCity catalog and six pinned CSV URLs → checksum-verified raw files',
                 'Spark CSV reader → selected columns → typed timestamps and PM2.5 → rejection audit',
                 'Duplicate resolution → hourly coverage filter → cached hourly DataFrame',
                 'Spark aggregations and SQL → percentile thresholds → timestamp joins → analytical tables',
                 'Matplotlib charts → independent pandas validation → report and handoff']:
        p(text)
    h('Implementation details',2)
    table(['File','Responsibility'],[
      ['src/data_ingestion.py','Download and verify manifest files; inspect all schemas and column completeness; load Spark CSVs.'],
      ['src/preprocessing.py','Typed parsing, study-window and quality rules, duplicate handling and hourly coverage.'],
      ['src/analysis.py','Station statistics and five temporal SQL queries.'],
      ['src/spike_detection.py','Exact percentile and strict spike indicators.'],
      ['src/synchronization.py','Pairwise overlap, independence expectation, lift, Jaccard and complete-case counts.'],
      ['src/visualization.py','Eight charts from exported tables.'],
      ['src/verification.py','Independent real-data recomputation and assertions.'],
      ['main.py','Execution, timed actions, output exports and validation.'],
      ['build_report.py','Generate this report and the group handoff from saved results.']])
    p('The implementation uses spark.read.csv, select, withColumn, try_cast, filter, dropDuplicates, groupBy, agg and join. SQL extracts local hours, weekdays and months. Caching materializes cleaned hourly data for repeated analyses. Every timed Spark transformation ends in an action such as count or collect, so the reported times include actual computation rather than lazy plan construction. Raw data is retained. Only the smaller processed and aggregate outputs are collected to the driver; this export design is suitable for the local dataset and would need revision for much larger outputs.')
    h('6 Results visualization and paper comparison')
    h('Results and visualizations',2)
    table(['Station','Hours','Mean','Median','Min','Max','SD'],[
      [short(r.station),f'{r.valid_hours:,}',f(r.mean),f(r.median),f(r.min),f(r.max),f(r.stddev)] for r in stats.itertuples()])
    p('All concentration summaries in the table are in µg/m³ and describe hourly means. SD is the sample standard deviation. Statistics use each station’s own available observations; they are not population exposure estimates or rankings of all Bengaluru neighborhoods.')
    hp=load('hourly_pattern');peak=hp.loc[hp['mean'].idxmax()]
    daily=load('daily');daily_good=daily[daily.n>=18];maxday=daily_good.loc[daily_good['mean'].idxmax()]
    weekend=load('weekday_weekend').pivot(index='station',columns='day_type',values='mean');diff=weekend.Weekend-weekend.Weekday
    monthly=load('monthly');monthrow=monthly.loc[monthly['mean'].idxmax()]
    var=stats.loc[stats.stddev.idxmax()]
    interpretations=[
      (f'{short(high.station)} records the highest observed station mean at {f(high["mean"])} µg/m³; {short(low.station)} records the lowest at {f(low["mean"])} µg/m³. Unequal missingness and station coverage can affect this ordering. On the all-six common hours, the highest mean is {short(common.iloc[0].station)} at {f(common.iloc[0]["mean"])} µg/m³.'),
      (f'The time series displays daily means only when at least 18 valid hourly means exist; unobserved days remain gaps. The largest displayed station-day mean is {f(maxday["mean"])} µg/m³ at {short(maxday.station)} on {maxday.day}. Individual peaks should be checked against instrument records before attributing a local cause.'),
      (f'The largest station-specific average hour profile is {f(peak["mean"])} µg/m³ at {short(peak.station)} in the {int(peak.hour_ist):02d}:00 IST hour. It averages available days and does not imply that every day peaks at this time.'),
      (f'Weekend minus weekday mean differences range from {f(diff.min())} to {f(diff.max())} µg/m³ across stations. The comparison is descriptive and unadjusted for season, weather, holidays or missing data; it does not identify a traffic effect.'),
      (f'The highest station-month mean is {f(monthrow["mean"])} µg/m³ at {short(monthrow.station)} in {monthrow.month}, based on {int(monthrow.n):,} valid hours. Month-to-month variation is visible, but the figure does not separate emissions from meteorology.'),
      (f'{short(var.station)} has the largest sample standard deviation at {f(var.stddev)} µg/m³. The IQR and coefficient of variation are also saved in station_statistics.csv so that absolute dispersion is not the only available comparison.'),
      ('The heatmap summarizes station-specific hourly patterns using the same means as the hour-profile chart. A common color scale supports comparison across stations. Missing source hours are excluded rather than colored as zero.'),
      (f'The strongest pair by lift is {short(top_pair.station_a)} and {short(top_pair.station_b)}: {int(top_pair.both_spike_hours):,} observed simultaneous spike hours versus {f(top_pair.expected_independent_hours)} under marginal independence, a lift of {f(top_pair.lift)}. The diagonal is omitted because it would compare a station with itself.')]
    for i,(name,interpretation) in enumerate(zip(m['charts'],interpretations),1):
        doc.add_page_break()
        h(f'Figure {i} '+['Station average','Pollution over time','Hour wise pattern','Weekday and weekend comparison','Monthly trend','Station variability','Station hourly heatmap','Spike synchronization heatmap'][i-1],2)
        doc.add_picture(str(ROOT/'outputs/charts'/name),width=Inches(6.55))
        p(interpretation);p('Source: generated Spark analytical tables from the six OpenCity station CSV files. Figure generated by src/visualization.py.')
        md.append(f'![Figure {i}](../outputs/charts/{name})\n')
    doc.add_page_break()
    h('Cross Station Synchronization Analysis',2)
    table(['Station','Threshold µg/m³','Spike hours','Valid hours'],[[short(r.station),f(r.threshold_p95),f'{r.spike_hours:,}',f'{r.valid_hours:,}'] for r in spikes.sort_values('station').itertuples()])
    p(f'There are {ncomplete:,} hours with valid observations at all six stations. Among these, {multiple:,} hours ({100*multiple/ncomplete:.2f}%) have at least two simultaneous station-relative spikes; {one:,} have exactly one. Among complete-case hours with any spike, the multiple-station share is {100*multiple/nspiking:.2f}%. An isolated flag refers to the monitored stations only and cannot establish that a pollution event is geographically isolated.')
    table(['Station pair','Overlap h','Both spike h','Expected h','Lift'],[[short(r.station_a)+' / '+short(r.station_b),f'{r.overlap_hours:,}',f'{r.both_spike_hours:,}',f(r.expected_independent_hours),f(r.lift)] for r in pairs.itertuples()])
    p('Pairwise overlap differs by pair. Thresholds are estimated from each station’s full valid study-window hours; the independence expectation uses spike marginals recomputed within each pair’s actual overlap. Positive lift may reflect common seasonal patterns or time dependence. No p-value, atmospheric trajectory or direction of movement is estimated.')
    h('Research Paper Comparison',2)
    table(['Dimension','Selected paper','Bengaluru AirWatch'],[
      ['Location','Ten northern Indian monitoring locations','Six stations within Bengaluru'],
      ['Variables','Six pollutants and meteorological variables','Hourly PM2.5 concentration'],
      ['Processing','R-based statistical and atmospheric analysis','Spark DataFrames and Spark SQL'],
      ['Question','Temporal patterns and meteorological associations','Station patterns and simultaneous relative spikes'],
      ['Transport evidence','Trajectory and directional analyses described','No transport model or source attribution'],
      ['Comparison boundary','Earlier period and different geography','Method-level comparison; numerical replication is not claimed']])
    p('The shared emphasis is temporal and spatial concentration variation. Our contribution is an auditable timestamp-aligned coincidence analysis with an explicit overlap denominator. The paper is used as an environmental analysis reference; it is not evidence that this local execution demonstrates Big Data scalability.')
    h('7 Performance evaluation and observations')
    h('Performance Evaluation',2)
    table(['Measure','Observed result'],[
      ['Input size',f'{m["raw_bytes"]:,} bytes'],['Input source rows',f'{m["raw_records"]:,}'],['Output station-hours',f'{m["output_hourly_records"]:,}'],
      ['Spark startup',f(m['startup_seconds'],3)+' s'],['CSV ingestion and cleaning actions',f(m['cleaning_seconds'],3)+' s'],
      ['Analytical aggregations and table materialization',f(m['aggregation_seconds'],3)+' s'],['Processed CSV export',f(m['export_seconds'],3)+' s'],
      ['Chart generation',f(m['visualization_seconds'],3)+' s'],['Independent verification',f(m['verification_seconds'],3)+' s'],
      ['Total local run excluding downloads',f(m['total_seconds_excluding_download'],3)+' s'],
      ['Warm Spark grouped statistics median',f(m['benchmark_median_seconds']['Spark'],4)+' s'],
      ['Warm pandas grouped statistics median',f(m['benchmark_median_seconds']['Pandas'],4)+' s']])
    p('The benchmark runs count, mean, minimum, maximum and sample standard deviation by station on the same hourly data. Both engines are warmed once and measured three times. Spark reads its cached hourly DataFrame; pandas uses an in-memory DataFrame exported from it. Parsing, cleaning, median calculation and startup are excluded from both benchmark timers. The three raw repetitions are saved. These timings describe this machine and workload only and do not compare complete end-to-end pipelines.')
    if m['benchmark_median_seconds']['Spark']>m['benchmark_median_seconds']['Pandas']:
        p('Pandas is faster for this small warm grouped aggregation in the measured local run. Spark scheduling and JVM coordination add overhead at this scale; no Spark speed advantage or cluster scalability is claimed.')
    else:p('Spark has the smaller measured median for this warm grouped aggregation. This narrow result does not establish a general speed advantage or cluster scalability.')
    table(['Test setup','Configuration'],[
      ['Operating system',m['platform']],['Processor',m['processor']],['Logical CPU count',m['logical_cpus']],
      ['Physical RAM',f(m['physical_memory_bytes']/1073741824,2)+' GiB'],['Python',m['python'].splitlines()[0]],
      ['PySpark',m['pyspark']],['Java',m['java_version'].splitlines()[0]],['Master','local[2]'],
      ['Shuffle partitions','4'],['Adaptive query execution','Enabled'],['Session time zone','Asia/Kolkata']])
    h('Observations and Discussion',2)
    p(f'The independent verification status is {check["status"]}. Checks compare every retained hourly row, station statistics, all exact thresholds, all spike counts, all pairwise overlaps and all temporal means and sample counts with pandas. Complete-case coincidence counts also agree. The common-hour station table helps assess the sensitivity of station rankings to missingness. The dataset supports evidence of coincident and single-station relative extremes, but does not explain their physical causes.')
    h('Limitations',2)
    for text in [
      'Only six selected monitoring locations are included, so the findings do not represent all Bengaluru neighborhoods or personal exposure.',
      'Coverage and missingness differ between stations. Complete-case restriction improves alignment but can itself select a nonrepresentative subset.',
      'The 15-minute source label does not guarantee independent instrument measurements; repeated concentrations can reflect source reporting or sensor averaging.',
      'Timestamp offsets are honored as encoded; an upstream timezone-labeling error would shift the hour-of-day interpretation.',
      'No sensor calibration or independent validation is available. Extreme valid measurements are retained and may influence statistics.',
      'Percentile thresholds are fitted on the analyzed period. The method is retrospective and does not evaluate forecasting or alarm accuracy.',
      'The independence baseline does not control for seasonality, meteorology or serial correlation. Lift is not a causal or inferential result.',
      'Local execution and driver-side exports do not establish cluster performance. The measured workload is small enough for pandas.'
    ]:bullet(text)
    h('Future Improvements',2)
    p('Confirm timestamp conventions and instrument quality flags with the publisher; expand station coverage; evaluate matched-season comparisons; use season-preserving shift or block-resampling baselines; and evaluate threshold stability on a held-out period. If substantially larger sources become available, write distributed Parquet outputs and measure multiple cluster sizes on the same workload.')
    h('Conclusion',2)
    p(f'The project executes a reproducible PM2.5 analysis over {m["raw_records"]:,} source rows and retains {m["output_hourly_records"]:,} valid station-hours. It produces eight charts, station and temporal summaries, all station-pair spike comparisons and measured execution evidence. The results show station differences and a mixture of single-station and synchronized high hours. They provide a basis for further investigation, with missingness, relative-threshold and local-performance limitations explicitly retained.')
    h('8 Tools responsible use and references')
    h('Tools and responsible use',2)
    p('Python, Apache Spark through PySpark, Spark SQL, pandas, NumPy and Matplotlib support execution and verification. python-docx produces the editable report; the rendering workflow exports its PDF. The original raw files remain unchanged and are excluded from Git. OpenCity catalog metadata labels the selected dataset Other (Public Domain); attribution to KSPCB, CPCB-labelled stations and OpenCity is retained. A manifest and downloader are provided instead of redistributing raw files in GitHub. The data are environmental station measurements, not individual personal records. Credentials are never committed. AI assisted code and report preparation; numerical findings are generated and verified by executable code.')
    h('References',2)
    for i,(name,text,url) in enumerate(REFERENCES,1):
        p(f'[{i}] {name}. {text} Access checked 31 October 2026.')
        hyperlink('Author supplied full text on ResearchGate' if i==4 else url,url)
    h('Appendix Source Code and Execution Instructions',2)
    hyperlink('GitHub repository','https://github.com/aganusha27/Bengaluru-AirWatch')
    p('The main branch contains source, generated aggregate tables and charts, metrics, report and learning material. Raw and processed row-level CSV files are excluded from Git but regenerate from the pinned manifest. See README.md for Windows and Unix setup and GROUP_HANDOFF.md for remaining submission details.')
    for cmd in ['python -m venv .venv','python -m pip install -r requirements.txt','python main.py','python build_report.py','python render_report.py','python verify_submission.py']:
        p(cmd)
    p('Activate the virtual environment before installation. Java 17 or 21 must be available to PySpark. In Windows PowerShell use .\\.venv\\Scripts\\Activate.ps1; in a Unix shell use source .venv/bin/activate. PDF rendering requires LibreOffice and Poppler; it is separate from the analytical run. If a source hash changes, inspect the new source before refreshing data/manifest.json. Do not bypass the checksum check silently.')
    doc.core_properties.title='Bengaluru AirWatch Air Pollution Analysis'
    doc.core_properties.subject='CSE557 empirical PM2.5 analysis'
    doc.core_properties.author=team.get('author','Bengaluru AirWatch Project Group')
    doc.save(ROOT/'report/Bengaluru_AirWatch_Final_Report.docx')
    (ROOT/'report/final_report.md').write_text('\n'.join(md),encoding='utf-8')
    (ROOT/'report/references.md').write_text('\n\n'.join(f'{name}. {text}\n{url}' for name,text,url in REFERENCES),encoding='utf-8')
    pd.DataFrame(evidence).to_csv(ROOT/'outputs/metrics/report_claims.csv',index=False)
    support_files(m,check,stats,spikes,pairs,ncomplete,multiple,team)
    print('Report DOCX and supporting files generated from verified outputs.')

def support_files(m,check,stats,spikes,pairs,ncomplete,multiple,team):
    def write(name,text):(ROOT/name).write_text(text,encoding='utf-8')
    setup='''```powershell
git clone https://github.com/aganusha27/Bengaluru-AirWatch.git
cd Bengaluru-AirWatch
python -m venv .venv
.\\.venv\\Scripts\\Activate.ps1
python -m pip install -r requirements.txt
python main.py
python build_report.py
python render_report.py
python verify_submission.py
```
Java 17 or 21 is required; Java 21 was used for the recorded run. Python 3.12 was used. On Linux/macOS, activate with `source .venv/bin/activate`. LibreOffice must be on PATH for `render_report.py`; alternatively pass `--soffice` with its absolute executable path. The analysis runs locally with two Spark worker threads. Permit local Java loopback communication if your sandbox requires execution outside its network isolation. No external service or API key is needed. To run analysis alone, stop after `python main.py`; the committed report is available for reading without installing a renderer.
'''
    write('README.md',f'''# Bengaluru AirWatch

Real-data PM2.5 analysis at six Bengaluru monitoring stations using PySpark 4.0.1.

## Recorded run
- {m['raw_records']:,} source rows in {m['raw_bytes']:,} bytes.
- {m['output_hourly_records']:,} valid station-hour observations.
- Eight charts and all 15 station-pair comparisons.
- Independent pandas verification: {check['status']}.
- Actual local total excluding downloads: {m['total_seconds_excluding_download']:.3f} seconds.

## Run
{setup}
The first run downloads six pinned OpenCity CSVs automatically and checks SHA-256 hashes. Raw files are retained locally, excluded from Git, and never replaced silently. To acquire only the files, use `python main.py --download-only`. Exact links and checksums are in `data/manifest.json`; catalog metadata are in `data/catalog_snapshot.json`.

## Outputs
`report/Bengaluru_AirWatch_Final_Report.docx` and `.pdf` contain the submission report. `outputs/tables` contains all generated CSV summaries; `outputs/charts` contains the eight figures. `outputs/metrics` records timing, configuration, source schema, checks and report evidence. `data/processed` is regenerated locally. `report/final_report.md` is a text version. The source PDFs and research full text are not redistributed.

## Method
PM2.5 concentration (µg/m³), not AQI. Honor encoded UTC timestamps, convert to IST and restrict to local 2024–2025. Reject absent/invalid/negative measurements with an audit; collapse identical selected readings; average conflicts; retain hourly means with >=3 valid distinct quarter-hours. Strictly exceed the station-specific exact 95th percentile for a spike. Compare station pairs only on actual overlapping hours. Independence expectation = marginal spike count A × count B / overlapping hours. Lift is descriptive, with no causal, transport, forecasting or significance claim.

## Performance and verification
`main.py` times actual Spark actions, not lazy transformation declarations. Warm grouped statistics use the same hourly data and equivalent count/mean/min/max/sample-SD operations in Spark and pandas, one warmup and three timed repetitions. Startup and cleaning are excluded from both benchmark timers. This small local workload does not establish cluster scalability. `src/verification.py` independently recomputes results from raw data using pandas; `verify_submission.py` validates delivered files and report evidence.

## Report regeneration
`python build_report.py` generates the DOCX and supporting Markdown after analysis. `python render_report.py --soffice /absolute/path/to/soffice --poppler /absolute/path/to/poppler/bin` renders the DOCX using an isolated LibreOffice profile and copies its PDF into `report/`. This optional step needs LibreOffice and Poppler installed from official sources. Analysis and source acquisition do not depend on those applications. Fill `report/team_details.json` using the example before rebuilding the final cover.

## Source and permissions
Dataset: https://data.opencity.in/dataset/bengaluru-hourly-air-quality-reports . Catalog license: Other (Public Domain). Attribute OpenCity, KSPCB and CPCB-labelled station sources. Access checked 31 October 2026. Raw and processed row-level files are excluded from Git; acquisition instructions and hashes allow reproduction.

## Research reference
Soni et al. (2024), Environmental Monitoring and Assessment 196:505. DOI https://doi.org/10.1007/s10661-024-12678-z . Publisher https://link.springer.com/article/10.1007/s10661-024-12678-z ; PubMed fallback https://pubmed.ncbi.nlm.nih.gov/38700603/ . This is a domain-method reference, not a Spark scalability study. See `report/references.md`.
''')
    write('data/README.md','''# Dataset

Six pinned 2024–2025 OpenCity resources: Bapuji Nagar, BTM Layout, Hebbal, Jayanagar 5th Block, Peenya and Silk Board. Download automatically with `python main.py --download-only`; files go in `data/raw/`.

If automated access is blocked, open the catalog URL in README.md, choose each matching resource named in `manifest.json`, click Download and save it under the exact `file` name in `data/raw/`. The manifest URL also opens the direct CSV. Rerun the command to check its hash. Do not use staging copies or replace a pinned dataset without inspection.

`manifest.json` supplies exact URLs, bytes, row counts, SHA-256 hashes and raw schema. `catalog_snapshot.json` supplies publisher and public-domain metadata. `outputs/tables/dataset_inventory.csv` supplies parsed coverage and schema compatibility; `column_completeness.csv` audits all fields. Source timestamps explicitly encode +0000; local-calendar calculations use IST. Selected schemas have concentrations rather than AQI. PM2.5, PM10, most gases and aromatics use µg/m³; CO uses mg/m³; NOx uses ppb; meteorological units are in headers.

Do not mix pollutants, fill missing measurements with zero, infer sensor timezone independently of metadata, or claim instrument validation. One one-field 2025 separator occurs in each file and is accounted for in the audit. The locally retained raw files are immutable; generated processed CSVs can be reproduced with `python main.py`.
''')
    pending='Add the group member names, USNs, responsibilities and faculty guide to report/team_details.json and regenerate the cover before submission.' if not team.get('members') else 'Verify group member and faculty details before submission.'
    write('GROUP_HANDOFF.md',f'''# Group handoff

## Overview
The empirical PM2.5 project is implemented and executed. Source rows: {m['raw_records']:,}; valid hourly records: {m['output_hourly_records']:,}. Eight charts and all 15 station pairs are generated. Independent verification passed. Actual metrics are in outputs/metrics/performance.json.

## Repository structure
- src/: ingestion, cleaning, summaries, spike detection, synchronization, charts and independent checks.
- data/: manifest, catalog snapshot and locally regenerated raw/processed folders.
- outputs/tables/: quality, temporal, station, coincidence and benchmark tables.
- outputs/charts/: all eight figures.
- outputs/metrics/: execution evidence and assertions.
- report/: final report, Markdown text, references and optional team details.
- main.py: pipeline; build_report.py: report; render_report.py: PDF; verify_submission.py: delivered-file checks.

## Completed tasks
Real OpenCity files acquired, hashes checked, actual schemas inspected, PySpark executed, all requested analyses completed, charts generated, outputs independently cross-checked, report generated, walkthrough and 25 viva questions prepared. Consult GITHUB_STATUS.json for verified push status and the verified artifact commit.

## Incomplete submission details
{pending}
The repository is private. A collaborator username was not supplied; no collaborator was added. The owner can open repository Settings → Collaborators → Add people, enter the intended GitHub username, and send the invitation. The friend must accept before cloning the private repository.

## Installation and execution
{setup}

## Known limitations
Unequal coverage, source timezone convention unconfirmed beyond encoded offsets, no sensor quality flags, retrospective thresholds, no autocorrelation-adjusted inference, no prediction or source attribution, and no distributed performance experiment. The selected paper is a domain analysis reference. See the final report for full qualifications.

## Important files and next actions
Read README.md, ASSIGNMENT_CHECKLIST.md, PROJECT_WALKTHROUGH.md, report/final_report.md and outputs/metrics/verification.json. Inspect raw source updates before changing pinned hashes. Rebuilding the report requires a completed local run; generating its PDF additionally requires LibreOffice and Poppler. Rerun checks after any numerical or methodological change. The recorded report and chart outputs can be reviewed without installing Spark.
''')
    required=[
      ('Real-world scenario and problem statement','1','src/analysis.py; report/final_report.md'),
      ('Recent research paper and existing approach','2','report/references.md; report/final_report.md'),
      ('Objectives and proposed analysis','3','src/preprocessing.py; src/spike_detection.py; src/synchronization.py'),
      ('Dataset identification and preparation','4','data/manifest.json; outputs/tables/dataset_inventory.csv; cleaning_audit.csv'),
      ('Big Data implementation and workflow','5','main.py; src/*.py; outputs/metrics/performance.json'),
      ('Results, visualization and paper comparison','6','outputs/charts/*.png; outputs/tables/*.csv; report/final_report.md'),
      ('Performance evaluation and observations','7','outputs/metrics/performance.json; outputs/tables/benchmark_repeats.csv; verification.json'),
      ('Tools, responsible use and references','8','requirements.txt; data/catalog_snapshot.json; report/references.md')]
    lines=['# Assignment checklist','','The eight mandatory faculty sections are addressed in the final report. Completion is supported by the executed outputs and independent checks, not just planned code. PDF and repository delivery are checked separately by verify_submission.py.','', '| Faculty requirement | Report section | Evidence | Status |','|---|---|---|---|']
    lines += [f'| {name} | {sec} | {files} | COMPLETE |' for name,sec,files in required]
    lines += ['','## Submission administration',pending,'Private collaborator invitation requires the collaborator username. This does not change the eight technical-content completion statuses. See GITHUB_STATUS.json for verified repository status.','', '## Deliberate analytical scope','No ML model, causal transport claim or distributed speedup is included. These are optional methods, not missing assignment outputs. The project uses the selected paper for temporal and spatial analysis context and explains the method differences.']
    write('ASSIGNMENT_CHECKLIST.md','\n'.join(lines))
    high=stats.iloc[0];pair=pairs.iloc[0]
    walkthrough=f'''# Project walkthrough

We study how PM2.5 varies across six Bengaluru monitoring stations and whether unusually high hours occur together. We chose this topic because the same city can contain different concentration patterns, and matching timestamps lets us distinguish coincident from single-station high readings.

## Actual dataset
The six real OpenCity files contain {m['raw_records']:,} source rows and matching schemas. Each row usually represents a station and a quarter-hour timestamp; six year separator rows are rejected. Station Name identifies the monitoring location; Timestamp encodes UTC; PM2.5 (µg/m³) measures fine-particle mass concentration. PM10, gases and meteorological columns are available but not combined into a single pollutant average. This project does not compute AQI.

## Why PySpark
Spark expresses a repeatable ingestion, cleaning, grouping and joining pipeline through DataFrames and SQL. It can execute these operations on distributed infrastructure, but our local run uses two worker threads and does not prove cluster scalability. Pandas is sufficient for this data size and independently checks the results.

## Workflow and files
data_ingestion.py checks pinned downloads and their schemas. preprocessing.py parses types, audits rejection reasons, deduplicates selected observations and averages at least three valid quarter-hours into each hourly mean. analysis.py computes station statistics and SQL time patterns. spike_detection.py creates exact 95th-percentile thresholds. synchronization.py joins overlapping hours and computes joint counts and lift. visualization.py makes eight charts from tables. verification.py separately recomputes the real-data results using pandas. main.py orchestrates actions and timing; build_report.py reads the saved evidence to create the report.

## Spark operations
read.csv reads the files; select chooses the measurement; withColumn and try_cast parse data; filter implements quality rules; dropDuplicates collapses repeated selected observations; groupBy and agg compute means, counts and percentiles; join aligns stations. SQL derives local hour, weekday and month. Caching avoids repeating the hourly calculation. count and collect force actual work before a timer stops.

## Spike detection
Sort the valid hourly means conceptually and choose the exact interpolated 95th percentile separately for each station. A spike strictly exceeds that station's value. It is unusually high for that station, not a health category or legal limit. Thresholds are fitted retrospectively, and ties can change the fraction flagged.

## Synchronization
For a station pair, inner-join timestamps so both observations exist. Count each station's spikes and simultaneous spikes on those matched hours. Expected simultaneous hours under independence = A × B / N. Lift = observed / expected. It describes coincidence and does not remove shared seasonality or prove transport. Across {ncomplete:,} complete six-station hours, {multiple:,} have at least two spikes.

## Actual results
Cleaning retains {m['output_hourly_records']:,} station-hour means. {str(high.station).split(',')[0]} has the highest available-hour mean, {high['mean']:.2f} µg/m³. The highest pairwise lift is {pair.lift:.2f}, for {str(pair.station_a).split(',')[0]} and {str(pair.station_b).split(',')[0]}. Read the full common-hour comparison and all pair denominators before interpreting these differences.

## The eight charts
1. Station average compares observed means, with unequal coverage.
2. Daily time series shows temporal change and gaps, requiring >=18 valid hours per plotted day.
3. Hour profile shows average local time-of-day patterns.
4. Weekday/weekend bars compare unadjusted hourly means.
5. Monthly lines show changing monthly means and differing coverage.
6. Standard deviation bars compare absolute variability; CSVs also include IQR and CV.
7. Station-hour heatmap displays the hour-profile matrix on a common color scale.
8. Spike lift heatmap compares pairwise simultaneous spikes with marginal independence.

## Performance
The actual total excluding downloads is {m['total_seconds_excluding_download']:.3f} seconds. Warm grouped-statistics medians are Spark {m['benchmark_median_seconds']['Spark']:.4f} s and pandas {m['benchmark_median_seconds']['Pandas']:.4f} s. Both benchmarks exclude startup, parsing and cleaning. They use the same hourly values and operations. The benchmark does not establish cluster scalability or compare complete pipelines.

## Research relationship
Soni et al. (2024) examine temporal and spatial pollutant patterns and meteorological associations in northern India. We study six stations within Bengaluru and compare timestamp-aligned relative spikes using Spark. We do not reproduce HYSPLIT or the paper's source-attribution methods. Its internally differing period descriptions are disclosed in the review.

## Limitations to explain to faculty
Coverage is incomplete; monitoring sites are not all neighborhoods; timezone conversion relies on source labels; sensor validation is unavailable; repeated source readings may not be independent; thresholds are retrospective; and lift is descriptive. No prediction model, causal transport analysis, or distributed speedup is claimed. The independent real-data checks in verification.json must pass before citing revised results.
'''
    write('PROJECT_WALKTHROUGH.md',walkthrough)
    qa=[
      ('What problem does the project solve?','It compares PM2.5 levels, timing and relative high hours across six Bengaluru stations and measures how often high hours coincide.'),
      ('Is this real time?','No. It analyzes historical batch data from resources labelled 15-minute observations. A live ingestion service is not implemented.'),
      ('What are the Big Data characteristics?','Volume is the saved source row and byte count; velocity refers to source reporting cadence; variety is the pollutant, weather and metadata fields; veracity covers missingness and measurement issues. Local size is modest.'),
      ('Why use Spark if pandas is sufficient?','The assignment requires a justified Big Data workflow. Spark implements distributed-style DataFrame operations; pandas is also used as an independent check. We do not claim Spark is necessary at this size.'),
      ('What is a Spark DataFrame?','A distributed table with named typed columns and an execution plan. Spark optimizes operations before actions execute them.'),
      ('What is lazy evaluation?','Transformations build a plan; actions such as count or collect run it. Timers must include an action to measure computation.'),
      ('Which operations are used?','CSV reading, column selection and parsing, filters, duplicate removal, grouped aggregations, exact percentiles and timestamp joins; SQL supplies calendar summaries.'),
      ('Where does the data come from?','Six OpenCity KSPCB/CPCB-labelled monitoring resources. Exact URLs, hashes, sizes and raw row counts are in data/manifest.json.'),
      ('Does the project analyze AQI?','No. The selected resources are labelled AQI but expose pollutant concentrations. We use PM2.5 in µg/m³ and do not convert it to AQI.'),
      ('How are timestamps handled?','We honor the encoded UTC offset, convert calendar calculations to Asia/Kolkata, and retain the local 2024–2025 window. The upstream sensor timezone convention remains a limitation.'),
      ('How are missing values handled?','Missing or invalid PM2.5 observations are excluded with reason counts. They are never replaced with zero. An hour needs three distinct valid quarter-hours.'),
      ('How are duplicates handled?','Identical selected station/time/PM2.5 rows collapse; conflicting values at one station and timestamp are averaged and counted. No arbitrary first-record choice is used.'),
      ('Why are hourly means used?','They align the supplied quarter-hour series at a common hourly interval while enforcing a minimum 75% within-hour observation count.'),
      ('How is a spike defined?','An hourly mean strictly above the station-specific exact 95th percentile of valid study-window hourly means. It is a retrospective statistical threshold.'),
      ('Are precisely 5% of readings spikes?','Not necessarily. Strict comparison, tied values and interpolated quantiles can change the flagged fraction.'),
      ('Is the threshold a health standard?','No. It describes the upper tail of that station’s observations and should not be interpreted as an AQI category or regulatory exceedance.'),
      ('How do stations align?','Pairs use an inner join on hourly timestamps. Missing observations cannot become zeros or non-spike flags. Six-station summaries use only complete-case hours.'),
      ('What does lift mean?','It is observed simultaneous spike hours divided by expected hours under independence, calculated using the same overlap-hour marginals. Values above one indicate more coincidence than that simple baseline.'),
      ('Does synchronization prove pollutant movement?','No. Weather, seasonal patterns and other shared factors can create coincidence. We do not estimate transport directions or causes.'),
      ('What are the main observed results?',f'{m["raw_records"]:,} source rows become {m["output_hourly_records"]:,} valid station-hours. There are {multiple:,} multiple-spike hours among {ncomplete:,} complete six-station hours. Exact station and pair results are in the tables.'),
      ('How is correctness checked?','A separate pandas implementation recomputes every processed hourly row, summary, percentile, spike count, pair overlap and temporal mean. The recorded checks passed.'),
      ('What performance comparison was measured?',f'Three warm grouped-statistic repetitions on the same hourly data: median Spark {m["benchmark_median_seconds"]["Spark"]:.4f} s, pandas {m["benchmark_median_seconds"]["Pandas"]:.4f} s. Startup and cleaning are excluded from both.'),
      ('Does the experiment establish scalability?','No. It runs locally on two worker threads with a modest dataset. Distributed hardware and increasing workload sizes would be required to test scalability.'),
      ('How does it differ from the research paper?','The paper studies multiple pollutants and meteorological associations across northern locations; this project examines Bengaluru PM2.5 and timestamp-aligned spike coincidence in Spark. It does not replicate trajectory modeling.'),
      ('What would you improve next?','Confirm timestamp and sensor metadata, add stations, evaluate held-out thresholds, control seasonality in the baseline, and measure distributed performance only if larger workloads justify it.')]
    write('FACULTY_VIVA_QUESTIONS.md','# Faculty viva questions\n\n'+'\n\n'.join(f'## {i} {q}\n\n{a}' for i,(q,a) in enumerate(qa,1)))
    write('report/team_details.example.json',json.dumps({'author':'Project group','faculty_guide':'','members':[{'name':'','usn':'','responsibility':''}]},indent=2))

if __name__=='__main__':build()

