"""Verify the delivered empirical report, chart links and required support files."""
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
import pandas as pd
from docx import Document
from pypdf import PdfReader

def main():
    root=Path(__file__).resolve().parent
    result=json.loads((root/'outputs/metrics/verification.json').read_text());assert result['status']=='PASS'
    perf=json.loads((root/'outputs/metrics/performance.json').read_text())
    manifest=json.loads((root/'data/manifest.json').read_text())
    assert sum(x['records'] for x in manifest)==perf['raw_records']
    assert sum(x['bytes'] for x in manifest)==perf['raw_bytes']
    for item in manifest:
        path=root/'data/raw'/item['file']
        if path.exists():assert hashlib.sha256(path.read_bytes()).hexdigest()==item['sha256']
    assert len(list((root/'outputs/charts').glob('*.png')))==8
    for f in ['README.md','GROUP_HANDOFF.md','ASSIGNMENT_CHECKLIST.md','PROJECT_WALKTHROUGH.md','FACULTY_VIVA_QUESTIONS.md']:
        assert (root/f).stat().st_size>500
    viva=(root/'FACULTY_VIVA_QUESTIONS.md').read_text(encoding='utf-8')
    assert len(re.findall(r'^## \d+ ',viva,re.M))==25
    checklist=(root/'ASSIGNMENT_CHECKLIST.md').read_text(encoding='utf-8')
    assert checklist.count('| COMPLETE |')==8
    doc=Document(root/'report/Bengaluru_AirWatch_Final_Report.docx')
    text='\n'.join(p.text for p in doc.paragraphs)
    assert len(doc.inline_shapes)==8
    for h in ['1 Real world scenario and problem statement','2 Recent research paper and existing approach',
              '3 Objectives and proposed analysis','4 Dataset identification and preparation',
              '5 Big Data implementation and workflow','6 Results visualization and paper comparison',
              '7 Performance evaluation and observations','8 Tools responsible use and references']:
        assert h in text
    assert f'{perf["raw_records"]:,}' in text and f'{perf["output_hourly_records"]:,}' in text
    assert f'{perf["total_seconds_excluding_download"]:,.3f}' in '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    stats=pd.read_csv(root/'outputs/tables/station_statistics.csv')
    assert f'{stats["mean"].max():,.2f}' in text
    claims=pd.read_csv(root/'outputs/metrics/report_claims.csv')
    expected=[perf['raw_records'],perf['output_hourly_records']]
    complete=pd.read_csv(root/'outputs/tables/complete_case_spike_distribution.csv')
    expected.extend([complete['count'].sum(),complete.loc[complete.spiking_stations>=2,'count'].sum(),
                     perf['cleaning_counts']['raw']-perf['cleaning_counts']['valid']])
    assert claims.value.tolist()==expected
    pdf=PdfReader(root/'report/Bengaluru_AirWatch_Final_Report.pdf')
    assert len(pdf.pages)>=10
    pdftext='\n'.join(p.extract_text() for p in pdf.pages)
    assert f'{perf["raw_records"]:,}' in pdftext and f'{perf["output_hourly_records"]:,}' in pdftext
    assert f'{perf["total_seconds_excluding_download"]:,.3f}' in pdftext, 'PDF is stale; rerun render_report.py after rebuilding the DOCX'
    assert not list(doc.styles['Title'].element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr'))
    pdf_links=[a.get_object().get('/A',{}).get('/URI') for p in pdf.pages for a in p.get('/Annots',[])]
    assert 'https://doi.org/10.1007/s10661-024-12678-z' in pdf_links
    assert 'https://data.opencity.in/dataset/bengaluru-hourly-air-quality-reports' in pdf_links
    assert not any(u and ('https://da'==u or 'www.data.opencity' in u) for u in pdf_links)
    output={'status':'PASS','faculty_sections':8,'charts':8,'viva_questions':25,'report_pages':len(pdf.pages),
            'report_images':len(doc.inline_shapes),'claim_ledger_values_checked':len(claims),
            'raw_hashes_verified_if_present':True,'independent_analytical_verification':'PASS',
            'note':'Numerical report text and tables are built from independently verified output tables; administrative team fields require user details.'}
    (root/'outputs/metrics/submission_verification.json').write_text(json.dumps(output,indent=2))
    print(json.dumps(output,indent=2))

if __name__=='__main__':main()
